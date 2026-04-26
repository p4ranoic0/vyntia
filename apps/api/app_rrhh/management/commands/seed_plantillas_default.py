# -*- coding: utf-8 -*-
"""
Comando de gestión: seed_plantillas_default

Genera las plantillas Word (.docx) por defecto para el sistema de documentos.
Crea los archivos .docx con marcadores {{VARIABLE}} y los registra en la BD
como PlantillaDocumento si aún no existen.

Uso:
    python manage.py seed_plantillas_default
    python manage.py seed_plantillas_default --force   # Recrear aunque existan
"""

import os
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand

from apps.documents.models import PlantillaDocumento


# ──────────────────────────────────────────────────────────────────────────────
# Helpers de formato python-docx
# ──────────────────────────────────────────────────────────────────────────────

def _set_font(run, name='Times New Roman', size_pt=11, bold=False):
    from docx.shared import Pt
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold


def _add_para(doc, text='', alignment='left', bold=False, size=11,
              space_before=0, space_after=6, font='Times New Roman'):
    """Agrega un párrafo con formato básico y devuelve el párrafo."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    ALIGN = {
        'left': WD_ALIGN_PARAGRAPH.LEFT,
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
        'justify': WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    para = doc.add_paragraph()
    para.alignment = ALIGN.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    if text:
        run = para.add_run(text)
        _set_font(run, name=font, size_pt=size, bold=bold)
    return para


def _add_run(para, text, bold=False, size=11, font='Times New Roman'):
    run = para.add_run(text)
    _set_font(run, name=font, size_pt=size, bold=bold)
    return run


def _firma_block(doc, etiqueta_izq='LA ENTIDAD', etiqueta_der='EL TRABAJADOR'):
    """Tabla de dos columnas para firmas."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc.add_paragraph()
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    # Fila 0: líneas de firma
    for i, txt in enumerate(['_' * 40, '_' * 40]):
        cell = table.cell(0, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(txt)
        _set_font(run, size_pt=11)
    # Fila 1: etiquetas
    for i, txt in enumerate([etiqueta_izq, etiqueta_der]):
        cell = table.cell(1, i)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(txt)
        _set_font(run, size_pt=10, bold=True)
    # Eliminar bordes de la tabla de firma
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)


# ──────────────────────────────────────────────────────────────────────────────
# Generadores de cada plantilla
# ──────────────────────────────────────────────────────────────────────────────

def _build_certificado_trabajo() -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    # Encabezado institucional
    _add_para(doc, '{{EMPRESA_NOMBRE}}', alignment='center', bold=True, size=13, space_after=2)
    _add_para(doc, 'RUC: {{EMPRESA_RUC}}  |  {{EMPRESA_DIRECCION}}',
              alignment='center', size=9, space_after=12)

    # Línea separadora
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('─' * 80)
    _set_font(run, size_pt=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Número de certificado
    _add_para(doc, 'CERTIFICADO DE TRABAJO N° {{NUMERO_CERTIFICADO}}',
              alignment='center', bold=True, size=14, space_before=6, space_after=16)

    # Firmante
    _add_para(doc,
              'EL JEFE DE LA OFICINA DE RECURSOS HUMANOS DE {{EMPRESA_NOMBRE}}, QUE SUSCRIBE, CERTIFICA:',
              alignment='justify', bold=True, size=11, space_after=10)

    # Cuerpo
    _add_para(doc,
              'Que, {{NOMBRE_COMPLETO}}, identificado(a) con {{TIPO_DOCUMENTO}} N° {{DNI}}, '
              'prestó servicios en {{EMPRESA_NOMBRE}} de acuerdo con el siguiente detalle:',
              alignment='justify', size=11, space_after=8)

    # Tabla de detalle
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Table Grid'
    filas = [
        ('Cargo:', '{{CARGO}}'),
        ('Área / Unidad Orgánica:', '{{AREA}}'),
        ('Fecha de ingreso:', '{{FECHA_INGRESO}}'),
        ('Régimen laboral:', 'Contrato Administrativo de Servicios (CAS) - D.L. N° 1057'),
    ]
    for i, (label, val) in enumerate(filas):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        r0 = c0.paragraphs[0].add_run(label)
        _set_font(r0, size_pt=11, bold=True)
        r1 = c1.paragraphs[0].add_run(val)
        _set_font(r1, size_pt=11)

    doc.add_paragraph()

    # Propósito opcional
    _add_para(doc,
              'El presente certificado se expide a solicitud del interesado(a) '
              'para fines de: {{PROPOSITO}}.',
              alignment='justify', size=11, space_after=8)

    _add_para(doc,
              'Se extiende el presente certificado para los fines que estime conveniente.',
              alignment='justify', size=11, space_after=16)

    # Fecha y ciudad
    _add_para(doc, '{{CIUDAD}}, {{FECHA_HOY}}',
              alignment='right', size=11, space_after=24)

    # Firma
    _firma_block(doc, 'JEFE DE RECURSOS HUMANOS', '{{NOMBRE_COMPLETO}}')

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_constancia_laboral() -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    _add_para(doc, '{{EMPRESA_NOMBRE}}', alignment='center', bold=True, size=13, space_after=2)
    _add_para(doc, 'RUC: {{EMPRESA_RUC}}  |  {{EMPRESA_DIRECCION}}',
              alignment='center', size=9, space_after=12)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('─' * 80)
    _set_font(run, size_pt=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_para(doc, 'CONSTANCIA DE TRABAJO',
              alignment='center', bold=True, size=14, space_before=6, space_after=16)

    _add_para(doc,
              'EL JEFE DE LA OFICINA DE RECURSOS HUMANOS DE {{EMPRESA_NOMBRE}}, QUE SUSCRIBE, DEJA CONSTANCIA:',
              alignment='justify', bold=True, size=11, space_after=10)

    _add_para(doc,
              'Que, {{NOMBRE_COMPLETO}}, identificado(a) con {{TIPO_DOCUMENTO}} N° {{DNI}}, '
              'presta servicios en {{EMPRESA_NOMBRE}}, desempeñándose como {{CARGO}} '
              'en el área de {{AREA}}, bajo la modalidad de Contrato Administrativo de Servicios (CAS) '
              'regulado por el Decreto Legislativo N° 1057, desde el {{FECHA_INGRESO}} hasta la fecha.',
              alignment='justify', size=11, space_after=12)

    _add_para(doc,
              'Se extiende la presente constancia a solicitud del interesado(a), '
              'para los fines que estime conveniente.',
              alignment='justify', size=11, space_after=20)

    _add_para(doc, '{{CIUDAD}}, {{FECHA_HOY}}',
              alignment='right', size=11, space_after=24)

    _firma_block(doc, 'JEFE DE RECURSOS HUMANOS', '{{NOMBRE_COMPLETO}}')

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_contrato() -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    _add_para(doc, '{{EMPRESA_NOMBRE}}', alignment='center', bold=True, size=13, space_after=2)
    _add_para(doc, 'RUC: {{EMPRESA_RUC}}  |  {{EMPRESA_DIRECCION}}',
              alignment='center', size=9, space_after=10)

    p = doc.add_paragraph()
    run = p.add_run('─' * 80)
    _set_font(run, size_pt=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)

    _add_para(doc, 'CONTRATO ADMINISTRATIVO DE SERVICIOS N° {{NUMERO_CONTRATO}}',
              alignment='center', bold=True, size=13, space_before=4, space_after=16)

    # Introducción
    _add_para(doc,
              'Conste por el presente documento el Contrato Administrativo de Servicios que celebran, '
              'de una parte, {{EMPRESA_NOMBRE}}, con RUC N° {{EMPRESA_RUC}}, con domicilio en '
              '{{EMPRESA_DIRECCION}}, representado por el Jefe de la Oficina de Administración, '
              'a quien en adelante se denominará LA ENTIDAD; y, de la otra parte '
              '{{NOMBRE_COMPLETO}}, identificado(a) con {{TIPO_DOCUMENTO}} N° {{DNI}}, '
              'a quien en adelante se denominará EL TRABAJADOR; en los términos y condiciones siguientes:',
              alignment='justify', size=11, space_after=12)

    clausulas = [
        ('CLÁUSULA PRIMERA: OBJETO DEL CONTRATO',
         'EL TRABAJADOR se desempeñará como {{CARGO}} en el área de {{AREA}}, '
         'cumpliendo las funciones descritas en la convocatoria correspondiente. '
         'Funciones específicas: {{FUNCIONES}}.'),
        ('CLÁUSULA SEGUNDA: PLAZO DEL CONTRATO',
         'La duración del presente Contrato se inicia el {{FECHA_INICIO}} y concluye '
         'el {{FECHA_FIN}}. El contrato podrá ser renovado según decisión de LA ENTIDAD.'),
        ('CLÁUSULA TERCERA: REMUNERACIÓN',
         'EL TRABAJADOR percibirá una remuneración mensual bruta de S/ {{SALARIO_BRUTO}} '
         '({{SALARIO_BRUTO}} soles), que incluye los montos y afiliaciones de ley, así como '
         'toda deducción aplicable.'),
        ('CLÁUSULA CUARTA: JORNADA Y HORARIO DE TRABAJO',
         'La jornada de trabajo es de {{JORNADA}}. El horario de prestación de servicios '
         'será: {{HORARIO_TRABAJO}}.'),
        ('CLÁUSULA QUINTA: LUGAR DE PRESTACIÓN DEL SERVICIO',
         'EL TRABAJADOR prestará los servicios en: {{LUGAR_TRABAJO}}. '
         'LA ENTIDAD podrá disponer la prestación de servicios fuera del lugar designado '
         'de acuerdo a las necesidades institucionales.'),
        ('CLÁUSULA SEXTA: OBLIGACIONES DEL TRABAJADOR',
         'Son obligaciones de EL TRABAJADOR: a) Cumplir las obligaciones derivadas del '
         'presente Contrato y las normas internas vigentes; b) Guardar reserva y '
         'confidencialidad respecto a la información institucional; c) No subcontratar '
         'total ni parcialmente la prestación de sus servicios.'),
        ('CLÁUSULA SÉPTIMA: DERECHOS DEL TRABAJADOR',
         'EL TRABAJADOR gozará de los derechos establecidos en el Decreto Legislativo '
         'N° 1057 y su Reglamento: descanso semanal, vacaciones de 30 días por año, '
         'prestaciones de salud (ESSALUD) y afiliación a un régimen de pensiones.'),
        ('CLÁUSULA OCTAVA: EXTINCIÓN DEL CONTRATO',
         'El contrato se extingue por: fallecimiento del trabajador, mutuo acuerdo, '
         'decisión unilateral debidamente justificada, o vencimiento del plazo pactado.'),
        ('CLÁUSULA NOVENA: DISPOSICIONES FINALES',
         'Las controversias derivadas del presente Contrato serán sometidas al Tribunal '
         'del Servicio Civil. En señal de conformidad, las partes suscriben el presente '
         'documento en dos ejemplares de igual validez, en {{CIUDAD}}, el {{FECHA_FIRMA}}.'),
    ]

    for titulo, cuerpo in clausulas:
        _add_para(doc, titulo, bold=True, size=11, space_before=6, space_after=2)
        _add_para(doc, cuerpo, alignment='justify', size=11, space_after=8)

    doc.add_paragraph()
    _firma_block(doc, 'LA ENTIDAD', 'EL TRABAJADOR')

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_adenda() -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    _add_para(doc, '{{EMPRESA_NOMBRE}}', alignment='center', bold=True, size=13, space_after=2)
    _add_para(doc, 'RUC: {{EMPRESA_RUC}}  |  {{EMPRESA_DIRECCION}}',
              alignment='center', size=9, space_after=10)

    p = doc.add_paragraph()
    run = p.add_run('─' * 80)
    _set_font(run, size_pt=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)

    _add_para(doc, 'ADENDA N° {{NUMERO_ADENDA}} AL CONTRATO N° {{NUMERO_CONTRATO}}',
              alignment='center', bold=True, size=13, space_before=4, space_after=16)

    _add_para(doc,
              'Conste por el presente documento la Adenda N° {{NUMERO_ADENDA}} al Contrato '
              'Administrativo de Servicios N° {{NUMERO_CONTRATO}}, que celebran de una parte '
              '{{EMPRESA_NOMBRE}}, con RUC N° {{EMPRESA_RUC}}, con domicilio en {{EMPRESA_DIRECCION}}, '
              'representado por el Jefe de la Oficina de Administración, a quien en adelante se '
              'denominará LA ENTIDAD; y, de la otra parte {{NOMBRE_COMPLETO}}, identificado(a) con '
              '{{TIPO_DOCUMENTO}} N° {{DNI}}, a quien en adelante se denominará EL TRABAJADOR.',
              alignment='justify', size=11, space_after=12)

    clausulas = [
        ('CLÁUSULA PRIMERA: ANTECEDENTES',
         'Con fecha {{FECHA_INICIO}}, LA ENTIDAD y EL TRABAJADOR suscribieron el Contrato '
         'Administrativo de Servicios N° {{NUMERO_CONTRATO}}, con el objeto de que EL TRABAJADOR '
         'preste servicios como {{CARGO}} en el área de {{AREA}}.'),
        ('CLÁUSULA SEGUNDA: MODIFICACIÓN',
         'Las partes acuerdan modificar las condiciones del contrato en los siguientes términos: '
         '{{FUNCIONES}}'),
        ('CLÁUSULA TERCERA: PLAZO',
         'Las modificaciones establecidas en la presente Adenda entran en vigencia a partir del '
         '{{FECHA_INICIO}} hasta el {{FECHA_FIN}}.'),
        ('CLÁUSULA CUARTA: REMUNERACIÓN',
         'La remuneración mensual bruta será de S/ {{SALARIO_BRUTO}} soles, incluyendo todos '
         'los beneficios y deducciones de ley correspondientes.'),
        ('CLÁUSULA QUINTA: CONDICIONES INALTERABLES',
         'Todas las demás cláusulas contractuales contenidas en el Contrato Administrativo de '
         'Servicios N° {{NUMERO_CONTRATO}}, así como las adendas suscritas con anterioridad, '
         'permanecen inalterables en lo que no se oponga a la presente Adenda.'),
        ('CLÁUSULA SEXTA: CONFORMIDAD',
         'En señal de conformidad y aprobación de las condiciones establecidas en el presente '
         'documento, LA ENTIDAD y EL TRABAJADOR lo suscriben en dos ejemplares igualmente '
         'válidos, en {{CIUDAD}}, el {{FECHA_FIRMA}}.'),
    ]

    for titulo, cuerpo in clausulas:
        _add_para(doc, titulo, bold=True, size=11, space_before=6, space_after=2)
        _add_para(doc, cuerpo, alignment='justify', size=11, space_after=8)

    doc.add_paragraph()
    _firma_block(doc, 'LA ENTIDAD', 'EL TRABAJADOR')

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ──────────────────────────────────────────────────────────────────────────────
# Definición de las plantillas a crear
# ──────────────────────────────────────────────────────────────────────────────

PLANTILLAS = [
    {
        'tipo': 'certificado_trabajo',
        'nombre': 'Certificado de Trabajo (Plantilla Base)',
        'descripcion': (
            'Plantilla base para certificados de trabajo. '
            'Incluye datos del empleado, cargo, área y período de servicio. '
            'Variables clave: {{NOMBRE_COMPLETO}}, {{DNI}}, {{CARGO}}, {{AREA}}, '
            '{{FECHA_INGRESO}}, {{NUMERO_CERTIFICADO}}, {{PROPOSITO}}, {{FECHA_HOY}}.'
        ),
        'builder': _build_certificado_trabajo,
        'archivo_nombre': 'plantilla_certificado_trabajo_base.docx',
    },
    {
        'tipo': 'constancia_laboral',
        'nombre': 'Constancia Laboral (Plantilla Base)',
        'descripcion': (
            'Plantilla base para constancias de trabajo vigente. '
            'Certifica que el empleado presta servicios actualmente. '
            'Variables clave: {{NOMBRE_COMPLETO}}, {{DNI}}, {{CARGO}}, {{AREA}}, '
            '{{FECHA_INGRESO}}, {{FECHA_HOY}}, {{CIUDAD}}.'
        ),
        'builder': _build_constancia_laboral,
        'archivo_nombre': 'plantilla_constancia_laboral_base.docx',
    },
    {
        'tipo': 'contrato',
        'nombre': 'Contrato CAS (Plantilla Base)',
        'descripcion': (
            'Plantilla base para contratos CAS (Decreto Legislativo N° 1057). '
            'Incluye 9 cláusulas estándar. '
            'Variables clave: {{NUMERO_CONTRATO}}, {{NOMBRE_COMPLETO}}, {{DNI}}, '
            '{{CARGO}}, {{AREA}}, {{FECHA_INICIO}}, {{FECHA_FIN}}, '
            '{{SALARIO_BRUTO}}, {{JORNADA}}, {{HORARIO_TRABAJO}}, {{LUGAR_TRABAJO}}, '
            '{{FUNCIONES}}, {{FECHA_FIRMA}}.'
        ),
        'builder': _build_contrato,
        'archivo_nombre': 'plantilla_contrato_cas_base.docx',
    },
    {
        'tipo': 'adenda',
        'nombre': 'Adenda al Contrato CAS (Plantilla Base)',
        'descripcion': (
            'Plantilla base para adendas (prórrogas/modificaciones) de contrato CAS. '
            'Variables clave: {{NUMERO_ADENDA}}, {{NUMERO_CONTRATO}}, {{NOMBRE_COMPLETO}}, '
            '{{DNI}}, {{CARGO}}, {{AREA}}, {{FECHA_INICIO}}, {{FECHA_FIN}}, '
            '{{SALARIO_BRUTO}}, {{FUNCIONES}}, {{FECHA_FIRMA}}.'
        ),
        'builder': _build_adenda,
        'archivo_nombre': 'plantilla_adenda_cas_base.docx',
    },
]


class Command(BaseCommand):
    help = 'Genera y registra las plantillas Word (.docx) base en el sistema.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--force',
            action='store_true',
            help='Recrear plantillas aunque ya existan en la BD.',
        )

    def handle(self, *args, **options):
        force = options['force']
        creadas = 0
        omitidas = 0

        for cfg in PLANTILLAS:
            tipo = cfg['tipo']
            nombre = cfg['nombre']

            # Verificar si ya existe
            existe = PlantillaDocumento.objects.filter(
                tipo=tipo, nombre=nombre
            ).exists()

            if existe and not force:
                self.stdout.write(f'  Omitida (ya existe): {nombre}')
                omitidas += 1
                continue

            # Eliminar la anterior si --force
            if existe and force:
                PlantillaDocumento.objects.filter(tipo=tipo, nombre=nombre).delete()
                self.stdout.write(self.style.WARNING(f'  Eliminada versión anterior: {nombre}'))

            try:
                docx_bytes = cfg['builder']()
                plantilla = PlantillaDocumento(
                    tipo=tipo,
                    nombre=nombre,
                    descripcion=cfg['descripcion'],
                    activa=True,
                )
                plantilla.archivo.save(
                    cfg['archivo_nombre'],
                    ContentFile(docx_bytes),
                    save=True,
                )
                creadas += 1
                self.stdout.write(
                    self.style.SUCCESS(f'  Creada: {nombre}  ({cfg["archivo_nombre"]})')
                )
            except Exception as exc:
                self.stdout.write(
                    self.style.ERROR(f'  ERROR al crear "{nombre}": {exc}')
                )

        self.stdout.write(
            self.style.SUCCESS(
                f'\nPlantillas: {creadas} creadas, {omitidas} omitidas.'
            )
        )
